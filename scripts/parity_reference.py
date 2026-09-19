"""Reference invariants computed with openpyxl (the library that generated the workbook)."""
import json, sys
import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

path = sys.argv[1]
wb = openpyxl.load_workbook(path)

def fill_key(cell):
    f = cell.fill
    if f is None or f.fill_type != 'solid': return None
    c = f.fgColor
    if c is None: return None
    if c.type == 'rgb' and isinstance(c.rgb, str): return 'rgb:' + c.rgb[-6:].upper()
    if c.type == 'theme':
        t = '%.2f' % (c.tint or 0.0)
        return 'theme:%d:%s' % (c.theme, '0.00' if t == '-0.00' else t)
    if c.type == 'indexed': return 'indexed:%d' % c.indexed
    return None

out = {'sheetNames': wb.sheetnames, 'merges': {}, 'strings': {}, 'formulas': {}, 'ints': {}, 'solidFills': {}, 'fillKeyCounts': {}, 'probes': {}}
for ws in wb.worksheets:
    out['merges'][ws.title] = len(ws.merged_cells.ranges)
    s = f = n = filled = 0
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell): continue
            v = cell.value
            if isinstance(v, str):
                if v.startswith('='): f += 1
                elif v != '': s += 1
            elif isinstance(v, (int, float)) and not isinstance(v, bool): n += 1
            k = fill_key(cell)
            if k is not None:
                filled += 1
                if ws.title.isdigit():
                    out['fillKeyCounts'][k] = out['fillKeyCounts'].get(k, 0) + 1
    out['strings'][ws.title] = s; out['formulas'][ws.title] = f; out['ints'][ws.title] = n; out['solidFills'][ws.title] = filled

for ref in ['1997!C4','1997!B9','1997!C9','1997!C3','1997!P21','2016!L61','2009!Y87','2010!A120','1998!Q92','2017!E8','2017!F8','2017!A9','2005!A1','2010!A117','2010!I117','2010!D116','2003!C2','2003!C13','2019!C9','2019!D9','Yachts!A1','Science!A4','Tours!A4']:
    sheet, a1 = ref.split('!')
    cell = wb[sheet][a1]
    if isinstance(cell, MergedCell):
        out['probes'][ref] = {'v': None, 'formula': None, 'fill': None, 'slave': True}
        continue
    v = cell.value
    formula = None
    if isinstance(v, str) and v.startswith('='): formula, v = v[1:], None
    out['probes'][ref] = {'v': v, 'formula': formula, 'fill': fill_key(cell), 'slave': False}

out['fillKeyCounts'] = dict(sorted(out['fillKeyCounts'].items()))
json.dump(out, sys.stdout, indent=1, sort_keys=True)
